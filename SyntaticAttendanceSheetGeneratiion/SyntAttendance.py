import random
import sys
import tempfile
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter, column_index_from_string

try:
    from PIL import Image, ImageEnhance, ImageFilter
except ImportError:
    raise ImportError(
        "Pillow is not installed in this interpreter. Install it with: pip install pillow"
    )

# =========================================================
# CONFIG
# =========================================================
BASE_DIR = Path(__file__).resolve().parent

TEMPLATE_PATH = BASE_DIR / "AttendanceTemplate.xlsx"
SIGNATURES_ROOT = Path(
    r"C:\Users\alper\Desktop\CNG 491\DataSet\SignatureGPDSSyntheticSignaturesManuscriptsv\firmasSINTESISmanuscritas"
)
OUTPUT_PATH = BASE_DIR / "Attendance_filled.xlsx"

SHEET_NAME = "Sheet2 (3)"
START_ROW = 5
STUDENT_COUNT = 24
FOLDER_COUNT = 24
SIGNATURE_COLUMNS = "E:I"

ID_COLUMN = "B"
LAST_NAME_COLUMN = "C"
FIRST_NAME_COLUMN = "D"

SEED = 42
ONE_FOLDER_PER_STUDENT = True

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tif", ".tiff", ".webp"}

TARGET_ROW_HEIGHT = 30
TARGET_COL_WIDTH = 14

MIN_IMAGE_WIDTH = 80
MIN_IMAGE_HEIGHT = 30

FIRST_NAMES = [
    "James", "John", "Michael", "David", "Daniel", "Matthew", "Andrew", "Joseph", "Christopher", "Robert",
    "William", "Thomas", "Anthony", "Joshua", "Ryan", "Nicholas", "Benjamin", "Samuel", "Henry", "Jack",
    "Alexander", "Ethan", "Lucas", "Noah", "Liam", "Oliver", "Nathan", "Adam", "Leo", "Connor",
    "Emma", "Olivia", "Sophia", "Isabella", "Mia", "Charlotte", "Amelia", "Harper", "Evelyn", "Abigail",
    "Emily", "Ella", "Grace", "Lily", "Chloe", "Hannah", "Scarlett", "Victoria", "Zoe", "Natalie"
]

LAST_NAMES = [
    "Smith", "Johnson", "Williams", "Brown", "Jones", "Miller", "Davis", "Garcia", "Rodriguez", "Wilson",
    "Martinez", "Anderson", "Taylor", "Thomas", "Moore", "Martin", "Jackson", "Thompson", "White", "Lopez",
    "Lee", "Gonzalez", "Harris", "Clark", "Lewis", "Young", "Allen", "King", "Wright", "Scott",
    "Green", "Baker", "Adams", "Nelson", "Hill", "Campbell", "Mitchell", "Roberts", "Carter", "Phillips"
]


def parse_columns(spec: str):
    spec = spec.strip().upper()
    if ":" in spec:
        start, end = spec.split(":", 1)
        start_i = column_index_from_string(start)
        end_i = column_index_from_string(end)
        if end_i < start_i:
            raise ValueError("Column range end cannot be before start.")
        return [get_column_letter(i) for i in range(start_i, end_i + 1)]
    return [part.strip() for part in spec.split(",") if part.strip()]


def list_signature_folders(root: Path):
    folders = [p for p in root.iterdir() if p.is_dir()]

    def folder_key(p: Path):
        try:
            return (0, int(p.name))
        except ValueError:
            return (1, p.name.lower())

    return sorted(folders, key=folder_key)


def is_image_usable(path: Path):
    try:
        with Image.open(path) as im:
            w, h = im.size
            return w >= MIN_IMAGE_WIDTH and h >= MIN_IMAGE_HEIGHT
    except Exception:
        return False


def find_images(root: Path):
    images = []
    for path in root.rglob("*"):
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS:
            if is_image_usable(path):
                images.append(path)
    return sorted(images)


def generate_unique_student_ids(count: int, rng: random.Random):
    ids = set()
    while len(ids) < count:
        ids.add("2" + "".join(str(rng.randint(0, 9)) for _ in range(6)))
    return list(ids)


def generate_unique_names(count: int, rng: random.Random):
    used = set()
    names = []
    max_unique = len(FIRST_NAMES) * len(LAST_NAMES)

    if count > max_unique:
        raise ValueError(f"Not enough unique name combinations for {count} students.")

    while len(names) < count:
        first = rng.choice(FIRST_NAMES)
        last = rng.choice(LAST_NAMES)
        key = (first, last)
        if key not in used:
            used.add(key)
            names.append((first, last))
    return names


def excel_col_width_to_pixels(width: float) -> int:
    return max(20, int((width or 8.43) * 7))


def excel_row_height_to_pixels(height_points: float) -> int:
    return max(20, int((height_points or 15) * 96 / 72))


def prepare_image_for_excel(src_path: Path, temp_dir: Path, row: int, col: str, target_w: int, target_h: int):
    with Image.open(src_path) as im:
        im = im.convert("L")
        im.thumbnail((target_w, target_h), Image.Resampling.LANCZOS)

        canvas = Image.new("L", (target_w, target_h), 255)
        x = (target_w - im.width) // 2
        y = (target_h - im.height) // 2
        canvas.paste(im, (x, y))

        canvas = ImageEnhance.Contrast(canvas).enhance(1.08)
        canvas = canvas.filter(ImageFilter.SHARPEN)

        out_path = temp_dir / f"sig_r{row}_{col}_{src_path.stem}.png"
        canvas.save(out_path, format="PNG")
        return out_path


def main():
    print("PYTHON:", sys.executable)
    print("Template path:", TEMPLATE_PATH)
    print("Template exists:", TEMPLATE_PATH.exists())
    print("Signatures root exists:", SIGNATURES_ROOT.exists())

    if not TEMPLATE_PATH.exists():
        raise FileNotFoundError(
            f"Template file not found: {TEMPLATE_PATH}\n"
            f"Put AttendanceTemplate.xlsx in the same folder as main.py"
        )

    if not SIGNATURES_ROOT.exists():
        raise FileNotFoundError(f"Signature root folder not found: {SIGNATURES_ROOT}")

    wb = load_workbook(TEMPLATE_PATH)

    if SHEET_NAME in wb.sheetnames:
        ws = wb[SHEET_NAME]
    else:
        ws = wb[wb.sheetnames[0]]
        print(f'Warning: "{SHEET_NAME}" not found. Using first sheet: "{ws.title}"')

    rng = random.Random(SEED)
    columns = parse_columns(SIGNATURE_COLUMNS)
    rows = list(range(START_ROW, START_ROW + STUDENT_COUNT))

    for col in columns:
        ws.column_dimensions[col].width = TARGET_COL_WIDTH

    for row in rows:
        ws.row_dimensions[row].height = TARGET_ROW_HEIGHT

    folders = list_signature_folders(SIGNATURES_ROOT)
    if not folders:
        raise FileNotFoundError(f"No subfolders found under: {SIGNATURES_ROOT}")

    selected_folders = folders[:FOLDER_COUNT]

    folder_to_images = {}
    for folder in selected_folders:
        imgs = find_images(folder)
        if imgs:
            folder_to_images[folder] = imgs

    if not folder_to_images:
        raise FileNotFoundError(
            "No usable signature images found in the selected folders. "
            "Try lowering MIN_IMAGE_WIDTH / MIN_IMAGE_HEIGHT."
        )

    usable_folders = list(folder_to_images.keys())

    if ONE_FOLDER_PER_STUDENT and len(usable_folders) < STUDENT_COUNT:
        raise ValueError(
            f"Not enough folders with usable images. Need {STUDENT_COUNT}, found {len(usable_folders)}."
        )

    student_ids = generate_unique_student_ids(STUDENT_COUNT, rng)
    student_names = generate_unique_names(STUDENT_COUNT, rng)

    if ONE_FOLDER_PER_STUDENT:
        assigned_folders = usable_folders[:STUDENT_COUNT]
    else:
        assigned_folders = [rng.choice(usable_folders) for _ in rows]

    with tempfile.TemporaryDirectory() as tmp:
        temp_dir = Path(tmp)

        for idx, row in enumerate(rows):
            student_id = student_ids[idx]
            first_name, last_name = student_names[idx]
            folder = assigned_folders[idx]
            images = folder_to_images[folder]

            ws[f"{ID_COLUMN}{row}"] = student_id
            ws[f"{LAST_NAME_COLUMN}{row}"] = last_name
            ws[f"{FIRST_NAME_COLUMN}{row}"] = f"{first_name}_synt"

            row_height = ws.row_dimensions[row].height or TARGET_ROW_HEIGHT

            for col in columns:
                img_path = rng.choice(images)
                cell = f"{col}{row}"
                col_width = ws.column_dimensions[col].width or TARGET_COL_WIDTH

                cell_h_px = excel_row_height_to_pixels(row_height)
                cell_w_px = excel_col_width_to_pixels(col_width)

                target_w = max(20, cell_w_px - 2)
                target_h = max(20, cell_h_px - 2)

                processed_img_path = prepare_image_for_excel(
                    img_path, temp_dir, row, col, target_w, target_h
                )

                img = XLImage(str(processed_img_path))
                img.anchor = cell
                ws.add_image(img)

            print(
                f"Filled row {row}: {student_id} - {first_name}_synt {last_name} - folder {folder.name}"
            )

        OUTPUT_PATH.parent.mkdir(parents=True, exist_ok=True)
        wb.save(OUTPUT_PATH)

    print("\nDone.")
    print(f"Output file: {OUTPUT_PATH}")
    print(f"Sheet used: {ws.title}")
    print(f"Rows filled: {rows[0]} - {rows[-1]}")
    print(f"Signature columns: {', '.join(columns)}")


if __name__ == "__main__":
    main()
